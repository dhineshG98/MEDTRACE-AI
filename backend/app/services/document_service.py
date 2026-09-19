"""Document orchestration.

Dispatches an uploaded file to the right extractor and returns a single
`ExtractionResult`. This is where the PDF -> OCR fallback decision is made.

Later phases hook AI classification and entity extraction onto the end of
`extract_text`; nothing above this layer needs to know how the text was read.
"""

from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional

from app.services import ocr_service, pdf_service
from app.utils.file_utils import IMAGE_EXTENSIONS
from app.utils.logger import get_logger, safe_preview

logger = get_logger(__name__)

try:
    import docx  # python-docx

    _DOCX_OK = True
except ImportError:  # pragma: no cover
    docx = None  # type: ignore[assignment]
    _DOCX_OK = False


class ExtractionError(Exception):
    """Text could not be extracted. Message is safe to show a user."""

    def __init__(self, code: str, message: str):
        self.code = code
        self.message = message
        super().__init__(message)


@dataclass
class ExtractionResult:
    text: str
    extraction_method: str  # pdf_text | ocr | docx | image_ocr
    page_count: Optional[int] = None
    ocr_pages: Optional[int] = None
    warnings: list[str] = field(default_factory=list)

    @property
    def char_count(self) -> int:
        return len(self.text)

    @property
    def word_count(self) -> int:
        return len(self.text.split())


def _extract_docx(path: Path) -> ExtractionResult:
    if not _DOCX_OK:
        raise ExtractionError("docx_unavailable", "DOCX support is not installed on the server.")
    try:
        document = docx.Document(str(path))
    except Exception as exc:  # noqa: BLE001
        logger.error("DOCX read failed: %s", type(exc).__name__)
        raise ExtractionError("docx_read_error", "The DOCX file could not be opened.") from exc

    parts = [p.text.strip() for p in document.paragraphs if p.text.strip()]

    # Medical DOCX files put lab values and medication lists in tables, so
    # tables are extracted as well, not just paragraphs.
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    text = "\n".join(parts)
    if not text.strip():
        raise ExtractionError("no_text_found", "No readable text was found in the document.")
    return ExtractionResult(text=text, extraction_method="docx")


def _extract_image(path: Path) -> ExtractionResult:
    available, detail = ocr_service.is_available()
    if not available:
        raise ExtractionError(
            "ocr_unavailable",
            "OCR is not available on the server, so image documents cannot be read.",
        )
    try:
        text = ocr_service.image_file_to_text(path)
    except ocr_service.OcrFailedError as exc:
        raise ExtractionError("ocr_failed", str(exc)) from exc

    if not text.strip():
        raise ExtractionError("no_text_found", "No readable text was found in the image.")
    return ExtractionResult(text=text, extraction_method="image_ocr", page_count=1, ocr_pages=1)


def _extract_pdf(path: Path) -> ExtractionResult:
    try:
        text, page_count = pdf_service.extract_text(path)
    except pdf_service.PdfUnavailableError as exc:
        raise ExtractionError("pdf_unavailable", "PDF support is not installed on the server.") from exc
    except pdf_service.PdfReadError as exc:
        raise ExtractionError("pdf_read_error", str(exc)) from exc

    if not pdf_service.needs_ocr(text):
        return ExtractionResult(text=text, extraction_method="pdf_text", page_count=page_count)

    # No usable text layer -> scanned document -> OCR fallback.
    logger.info("PDF has no usable text layer; falling back to OCR")
    available, detail = ocr_service.is_available()
    if not available:
        raise ExtractionError(
            "ocr_unavailable",
            "This looks like a scanned PDF, but OCR is not available on the server.",
        )

    try:
        images = pdf_service.render_pages_to_images(path)
        ocr_text, ocr_pages = ocr_service.images_to_text(images)
    except pdf_service.PdfReadError as exc:
        raise ExtractionError("pdf_read_error", str(exc)) from exc
    except ocr_service.OcrFailedError as exc:
        raise ExtractionError("ocr_failed", str(exc)) from exc

    warnings: list[str] = []
    if page_count and ocr_pages < page_count:
        warnings.append(f"OCR produced text for {ocr_pages} of {page_count} pages.")

    return ExtractionResult(
        text=ocr_text,
        extraction_method="ocr",
        page_count=page_count,
        ocr_pages=ocr_pages,
        warnings=warnings,
    )


def _run_extract(path: Path, extension: str) -> ExtractionResult:
    ext = extension.lower().lstrip(".")

    if ext == "pdf":
        result = _extract_pdf(path)
    elif ext == "docx":
        result = _extract_docx(path)
    elif ext in IMAGE_EXTENSIONS:
        result = _extract_image(path)
    else:
        raise ExtractionError("unsupported_file_type", f"'{ext.upper()}' cannot be processed.")

    logger.info(
        "Extracted via %s: %d chars, %d words | %s",
        result.extraction_method,
        result.char_count,
        result.word_count,
        safe_preview(result.text),
    )
    return result


def extract_text(path: Path, extension: str) -> ExtractionResult:
    """Extract text from a stored document. Raises ExtractionError on failure.

    If the document is encrypted on disk (.enc), it is decrypted ephemerally into
    a transient memory-backed or shred-unlinked file, processed, and immediately destroyed.
    """
    from app.services.encryption_service import DecryptionError, ephemeral_decrypted_file

    if path.suffix == ".enc":
        try:
            with ephemeral_decrypted_file(path, extension) as temp_plaintext_path:
                return _run_extract(temp_plaintext_path, extension)
        except DecryptionError as exc:
            logger.error("Failed to decrypt document at %s: %s", path, exc)
            raise ExtractionError("decryption_failed", "Failed to decrypt stored medical document.") from exc
    return _run_extract(path, extension)


def capabilities() -> dict[str, tuple[bool, str]]:
    """Reports which extractors are usable, for the health endpoint."""
    pdf_ok, pdf_detail = pdf_service.is_available()
    ocr_ok, ocr_detail = ocr_service.is_available()
    return {
        "pdf": (pdf_ok, pdf_detail),
        "ocr": (ocr_ok, ocr_detail),
        "docx": (_DOCX_OK, "python-docx ready" if _DOCX_OK else "python-docx not installed"),
    }
